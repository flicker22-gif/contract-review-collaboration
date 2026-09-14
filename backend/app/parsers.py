"""把上传的 docx / pdf 解析成有序段落列表，供前端渲染和批注锚定。

输入为文件字节（而非路径），与存储后端解耦：本地磁盘、对象存储
取回的字节走同一条解析路径。
"""
import io
import re
from typing import List


def parse_docx(data: bytes) -> List[str]:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    paragraphs: List[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # 表格内容也按行提取，避免表格里的条款丢失
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    return paragraphs


def parse_pdf(data: bytes) -> List[str]:
    import fitz  # PyMuPDF

    paragraphs: List[str] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            for block in page.get_text("blocks"):
                # block: (x0, y0, x1, y1, text, block_no, block_type)
                if len(block) < 5 or block[6] != 0:  # 只要文本块
                    continue
                text = re.sub(r"\s*\n\s*", " ", block[4]).strip()
                if text:
                    paragraphs.append(text)
    return paragraphs


def parse_document(data: bytes, file_type: str) -> List[str]:
    if file_type == "docx":
        return parse_docx(data)
    if file_type == "pdf":
        return parse_pdf(data)
    raise ValueError(f"不支持的文件类型: {file_type}")
